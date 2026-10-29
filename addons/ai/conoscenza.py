"""La conoscenza dell'assistente: indice dei testi (record, note del chatter, documenti) e ricerca ibrida.

- Indice per azienda, permessi alla ricerca: un risultato si mostra solo se la ModelView del record d'origine lo
  fa vedere all'utente (anche le cartelle dei Documenti, che filtrano nella loro vista).
- Vettori di EmbeddingGemma 2 accorciati a 256 dimensioni (float32 nel database) dal servizio indicato in
  `ai.url_embeddings`; confronto in Python. Senza servizio resta la ricerca per parole chiave.
- L'indicizzazione è incrementale: ogni frammento ha un'impronta del testo, si ricalcola solo se cambia; i
  frammenti di record, note o documenti spariti si tolgono.
"""
import hashlib
import io
import json
import logging
import re
import time
import urllib.request
from array import array

from django.conf import settings
from django.contrib.contenttypes.models import ContentType
from django.db import models

from .models import Frammento

log = logging.getLogger("ai.conoscenza")

DIMENSIONI = 256
LUNGHEZZA = 1200                      # caratteri per frammento di documento
SOVRAPPOSIZIONE = 150
MAX_TESTO_RECORD = 2000
MAX_TESTO_DOCUMENTO = 400_000         # oltre, il documento si indicizza solo in parte
LOTTO = 64                            # testi per chiamata al servizio di embedding
TIMEOUT = 60
ESCLUSI = {"ai"}
# viste tecniche del kernel: coda dei lavori, email grezze, profili e definizioni dei campi (rumore o riservate)
MODELLI_ESCLUSI = {"core.lavoro", "core.emailricevuta", "core.profiloutente", "core.campopersonalizzato"}
SENSIBILE = re.compile(r"password|chiave|token|segret|secret|cifrat|hash|iban", re.IGNORECASE)
PAROLA = re.compile(r"\w{3,}", re.UNICODE)
PAROLE_VUOTE = {"che", "per", "con", "del", "della", "delle", "dei", "degli", "dal", "dalla", "nel", "nella", "sul",
                "sulla", "una", "uno", "gli", "le", "non", "come", "cosa", "dice", "sono", "questo", "questa", "quali",
                "quale", "quanto", "tutti", "tutte", "anche", "più", "mio", "mia", "miei", "loro", "the", "and"}


# ---------- testo dai file ----------

def testo_da_file(nome, dati):
    """Il testo di un file (txt, md, csv, PDF, Word, Excel); "" se il tipo non è supportato o il file è illeggibile."""
    estensione = nome.lower().rsplit(".", 1)[-1] if "." in nome else ""
    try:
        if estensione in ("txt", "md", "csv", "tsv", "json", "xml", "html", "htm", "eml"):
            testo = dati.decode("utf-8", errors="replace")
            return re.sub(r"<[^>]+>", " ", testo) if estensione in ("html", "htm", "xml") else testo
        if estensione == "pdf":
            from pypdf import PdfReader
            return "\n\n".join((pagina.extract_text() or "") for pagina in PdfReader(io.BytesIO(dati)).pages)
        if estensione == "docx":
            import docx
            documento = docx.Document(io.BytesIO(dati))
            righe = [p.text for p in documento.paragraphs]
            for tabella in documento.tables:
                righe += [" | ".join(c.text for c in riga.cells) for riga in tabella.rows]
            return "\n".join(righe)
        if estensione in ("xlsx", "xlsm"):
            return "\n".join(" | ".join(map(str, riga)) for _, righe in fogli(dati) for riga in righe)
    except Exception:                                  # file rotto o protetto: non si indicizza, non si blocca
        log.warning("testo non estratto da %s", nome, exc_info=True)
    return ""


def fogli(dati, massimo=5000):
    """[(nome del foglio, [[valori]])] di un file Excel, in sola lettura."""
    from openpyxl import load_workbook
    libro = load_workbook(io.BytesIO(dati), read_only=True, data_only=True)
    risultato = []
    for foglio in libro.worksheets:
        righe = []
        for riga in foglio.iter_rows(values_only=True):
            if any(v not in (None, "") for v in riga):
                righe.append(["" if v is None else v for v in riga])
            if len(righe) >= massimo:
                break
        risultato.append((foglio.title, righe))
    libro.close()
    return risultato


def pezzi(testo, lunghezza=LUNGHEZZA, sovrapposizione=SOVRAPPOSIZIONE):
    """Il testo in pezzi di circa `lunghezza` caratteri, tagliati a fine paragrafo o frase quando si può."""
    testo = re.sub(r"[ \t]+", " ", re.sub(r"\n{3,}", "\n\n", testo or "")).strip()
    risultato, inizio = [], 0
    while inizio < len(testo):
        fine = min(len(testo), inizio + lunghezza)
        if fine < len(testo):
            taglio = max(testo.rfind("\n\n", inizio + lunghezza // 2, fine), testo.rfind(". ", inizio + lunghezza // 2, fine))
            if taglio > inizio:
                fine = taglio + 1
        risultato.append(testo[inizio:fine].strip())
        if fine >= len(testo):
            break
        inizio = max(fine - sovrapposizione, inizio + 1)
    return [p for p in risultato if p]


# ---------- sorgenti ----------

def _impronta(testo):
    return hashlib.sha1(testo.encode()).hexdigest()


def _viste_indicizzabili():
    from core.viste import registry
    registry.autodiscover(["core", *settings.ADDONS_ATTIVI])
    for modulo in ["core", *settings.ADDONS_ATTIVI]:
        if modulo in ESCLUSI:
            continue
        for vista in registry.viste_di(modulo):
            if vista.opts.label_lower in MODELLI_ESCLUSI:
                continue
            campi = [f for f in vista.opts.concrete_fields
                     if isinstance(f, (models.CharField, models.TextField)) and not f.choices
                     and not SENSIBILE.search(f.name)]
            if campi and not (modulo == "documenti" and vista.model.__name__ == "Documento"):
                yield vista, campi


def frammenti_record():
    """(chiave, dati del frammento) per ogni record delle ModelView con campi di testo."""
    for vista, campi in _viste_indicizzabili():
        ct = ContentType.objects.get_for_model(vista.model)
        azienda_campo = vista.company_field
        for obj in vista.model._default_manager.all().iterator(chunk_size=500):
            righe = [f"{vista.verbose_name}: {obj}"]
            for f in campi:
                valore = getattr(obj, f.attname)
                if valore and str(valore) != str(obj):
                    righe.append(f"{f.verbose_name}: {valore}")
            testo = "\n".join(righe)[:MAX_TESTO_RECORD]
            yield {"fonte": Frammento.RECORD, "content_type": ct, "object_id": obj.pk, "chiave": "",
                   "azienda_id": getattr(obj, f"{azienda_campo}_id", None) if azienda_campo else None,
                   "titolo": f"{str(vista.verbose_name).capitalize()} «{obj}»"[:250], "testo": testo}


def frammenti_note():
    """Note e messaggi del chatter (non i tracciamenti né gli eventi di sistema), legati al loro record."""
    from core.models import Messaggio
    for m in (Messaggio.objects.filter(tipo__in=[Messaggio.NOTA, Messaggio.MESSAGGIO]).exclude(corpo="")
              .select_related("content_type", "autore").iterator(chunk_size=500)):
        record = m.record
        if record is None:
            continue
        autore = (m.autore.get_full_name() or m.autore.username) if m.autore else (m.autore_email or "esterno")
        yield {"fonte": Frammento.NOTA, "content_type": m.content_type, "object_id": m.object_id,
               "chiave": str(m.pk), "azienda_id": getattr(record, "azienda_id", None),
               "titolo": f"{m.get_tipo_display()} di {autore} del {m.creato:%d/%m/%Y} su «{record}»"[:250],
               "testo": m.corpo[:MAX_TESTO_RECORD]}


def frammenti_documenti():
    if "documenti" not in settings.ADDONS_ATTIVI:
        return
    from documenti.models import Documento
    ct = ContentType.objects.get_for_model(Documento)
    for d in Documento.objects.filter(cestinato_il__isnull=True).iterator(chunk_size=200):
        file = d.contenuto
        if not file:
            continue
        try:
            with file.open("rb") as f:
                dati = f.read(20 * 1024 * 1024)
        except (OSError, ValueError):
            continue
        testo = testo_da_file(d.nome, dati)[:MAX_TESTO_DOCUMENTO]
        if d.descrizione:
            testo = f"{d.descrizione}\n\n{testo}"
        for i, pezzo in enumerate(pezzi(testo)):
            yield {"fonte": Frammento.DOCUMENTO, "content_type": ct, "object_id": d.pk, "chiave": str(i),
                   "azienda_id": d.azienda_id, "titolo": f"Documento «{d.nome}»" + (f", parte {i + 1}" if i else ""),
                   "testo": pezzo}


SORGENTI = {Frammento.RECORD: frammenti_record, Frammento.NOTA: frammenti_note,
            Frammento.DOCUMENTO: frammenti_documenti}


# ---------- embedding ----------

def url_servizio():
    from core import parametri
    return (parametri.leggi("ai.url_embeddings") or "").strip().rstrip("/")


def embedding(testi, tipo="document"):
    """Vettori normalizzati (liste di float) dal servizio; None se il servizio non c'è o non risponde."""
    url = url_servizio()
    if not url or not testi:
        return None
    richiesta = urllib.request.Request(f"{url}/v1/embeddings", headers={"Content-Type": "application/json"},
                                       data=json.dumps({"input": testi, "input_type": tipo,
                                                        "dimensions": DIMENSIONI}).encode())
    try:
        with urllib.request.urlopen(richiesta, timeout=TIMEOUT) as risposta:
            dati = json.load(risposta)
        return [v["embedding"] for v in sorted(dati["data"], key=lambda v: v["index"])]
    except Exception as e:
        log.warning("servizio di embedding non disponibile (%s): %s", url, e)
        return None


def _binario(vettore):
    return array("f", vettore).tobytes()


def _da_binario(dati):
    v = array("f")
    v.frombytes(bytes(dati))
    return v


# ---------- indicizzazione ----------

def indicizza(fonti=None, limite_secondi=240):
    """Allinea l'indice alle sorgenti. Restituisce un riepilogo {nuovi, aggiornati, tolti, vettori}."""
    inizio = time.monotonic()
    conti = {"nuovi": 0, "aggiornati": 0, "tolti": 0, "vettori": 0}
    for fonte in fonti or list(SORGENTI):
        esistenti = {(f.content_type_id, f.object_id, f.chiave): f for f in
                     Frammento.objects.filter(fonte=fonte).only("pk", "content_type_id", "object_id", "chiave",
                                                                "impronta", "azienda_id", "titolo")}
        visti = set()
        for dati in SORGENTI[fonte]():
            chiave = (dati["content_type"].pk, dati["object_id"], dati["chiave"])
            visti.add(chiave)
            impronta = _impronta(dati["titolo"] + "\n" + dati["testo"])
            f = esistenti.get(chiave)
            if f is None:
                Frammento.objects.create(impronta=impronta, **dati)
                conti["nuovi"] += 1
            elif f.impronta != impronta or f.azienda_id != dati["azienda_id"]:
                Frammento.objects.filter(pk=f.pk).update(impronta=impronta, titolo=dati["titolo"], testo=dati["testo"],
                                                         azienda_id=dati["azienda_id"], vettore=None)
                conti["aggiornati"] += 1
        spariti = [f.pk for k, f in esistenti.items() if k not in visti]
        for i in range(0, len(spariti), 500):
            conti["tolti"] += Frammento.objects.filter(pk__in=spariti[i:i + 500]).delete()[0]
    while time.monotonic() - inizio < limite_secondi:             # vettori mancanti, a lotti
        da_fare = list(Frammento.objects.filter(vettore__isnull=True).only("pk", "titolo", "testo")[:LOTTO])
        if not da_fare:
            break
        vettori = embedding([f"{f.titolo}\n{f.testo}" for f in da_fare])
        if vettori is None:
            break
        for f, v in zip(da_fare, vettori):
            Frammento.objects.filter(pk=f.pk).update(vettore=_binario(v))
        conti["vettori"] += len(da_fare)
    return conti


def indicizza_cron():
    return json.dumps(indicizza(), ensure_ascii=False)


# ---------- ricerca ----------

def parole(testo):
    return {p for p in (w.lower() for w in PAROLA.findall(testo or "")) if p not in PAROLE_VUOTE}


def _coseno(a, b):
    return sum(x * y for x, y in zip(a, b))                  # vettori già normalizzati


def cerca(domanda, utente, azienda=None, fonti=None, app="", limite=8, visibile=None):
    """[(frammento, punteggio)] più pertinenti tra quelli che `visibile(frammento)` lascia vedere."""
    termini = parole(domanda)
    qs = Frammento.objects.all()
    qs = qs.filter(models.Q(azienda_id=getattr(azienda, "pk", None)) | models.Q(azienda__isnull=True))
    if fonti:
        qs = qs.filter(fonte__in=fonti)
    vettore = embedding([domanda], "query")
    domanda_v = array("f", vettore[0]) if vettore else None
    candidati = []
    for f in qs.only("pk", "fonte", "content_type_id", "object_id", "titolo", "testo", "vettore", "chiave").iterator(
            chunk_size=2000):
        testo_parole = parole(f.titolo + " " + f.testo)
        chiave = len(termini & testo_parole) / len(termini) if termini else 0.0
        if domanda_v is not None and f.vettore:
            punteggio = 0.8 * _coseno(domanda_v, _da_binario(f.vettore)) + 0.2 * chiave
        else:
            punteggio = chiave
        if punteggio > 0:
            candidati.append((punteggio, f))
    candidati.sort(key=lambda c: c[0], reverse=True)
    risultati, per_record = [], {}
    for punteggio, f in candidati[:limite * 8]:
        if app and _app_di(f) != app:
            continue
        if visibile is not None and not visibile(f):
            continue
        record = (f.content_type_id, f.object_id)
        if per_record.get(record, 0) >= 2:                   # al massimo due pezzi dello stesso documento o record
            continue
        per_record[record] = per_record.get(record, 0) + 1
        risultati.append((f, round(punteggio, 3)))
        if len(risultati) >= limite:
            break
    return risultati


def _app_di(frammento):
    from core.viste import registry
    modello = ContentType.objects.get_for_id(frammento.content_type_id).model_class()
    vista = registry.per_modello(modello) if modello else None
    return vista.modulo if vista else ""


def statistiche():
    totale = Frammento.objects.count()
    return {"frammenti": totale, "con_vettore": Frammento.objects.filter(vettore__isnull=False).count(),
            "servizio": url_servizio() or "non configurato"}

