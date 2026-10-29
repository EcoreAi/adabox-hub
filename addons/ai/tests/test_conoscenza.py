import io
import math
import re
import zlib

import pytest
from django.conf import settings
from django.contrib.auth.models import Group, Permission, User
from django.core.files.base import ContentFile
from django.core.management import call_command

from ai import conoscenza, motore, strumenti
from ai.gateway import finto
from ai.models import AnelloCatena, Conversazione, Evento, Frammento, Modello, Profilo, Provider
from ai.strumenti import conoscenza as strumenti_conoscenza
from core import chatter
from core.models import Azienda

con_documenti = pytest.mark.skipif("documenti" not in settings.ADDONS_ATTIVI, reason="serve l'app documenti")


def vettore_finto(testo):
    """Embedding finto ma "semantico": sinonimi nello stesso cassetto, vettore normalizzato a 256."""
    sinonimi = {"rinnovo": "rinnov", "rinnovato": "rinnov", "rinnova": "rinnov", "proroga": "rinnov",
                "disdetta": "disdett", "recesso": "disdett"}
    v = [0.0] * conoscenza.DIMENSIONI
    for parola in re.findall(r"\w+", testo.lower()):
        v[zlib.crc32(sinonimi.get(parola, parola).encode()) % conoscenza.DIMENSIONI] += 1   # stabile tra processi
    norma = math.sqrt(sum(x * x for x in v)) or 1
    return [x / norma for x in v]


@pytest.fixture
def con_servizio(monkeypatch):
    chiamate = []

    def embedding(testi, tipo="document"):
        chiamate.append((tipo, len(testi)))
        return [vettore_finto(t) for t in testi]
    monkeypatch.setattr(conoscenza, "embedding", embedding)
    monkeypatch.setattr(conoscenza, "url_servizio", lambda: "http://finto")
    return chiamate


@pytest.fixture
def scena(db):
    capo = User.objects.create_superuser("capo", "capo@example.com", "x")
    azienda = Azienda.objects.get(pk=1)
    Azienda.objects.filter(pk=1).update(nome="Alfa Spa")
    bianchi = Azienda.objects.create(nome="Bianchi Srl", citta="Milano")
    chatter.scrivi(bianchi, capo, "Il contratto di manutenzione si rinnova ogni anno a gennaio, salvo disdetta "
                                  "con 60 giorni di preavviso.", tipo="nota")
    chatter.scrivi(bianchi, capo, "Chiamato per la fattura di ottobre.", tipo="nota")
    return {"capo": capo, "azienda": azienda, "bianchi": bianchi}


def _contesto(utente, azienda):
    return strumenti.Contesto(utente=utente, azienda=azienda, conversazione=Conversazione.objects.create(utente=utente))


def test_pezzi_con_sovrapposizione():
    testo = " ".join(f"Frase numero {i} del contratto." for i in range(200))
    parti = conoscenza.pezzi(testo, lunghezza=300, sovrapposizione=50)
    assert len(parti) > 10 and all(len(p) <= 300 for p in parti)
    assert parti[0].endswith(".") and parti[1][:20] in parti[0][-80:]               # tagli a fine frase, sovrapposti
    assert conoscenza.pezzi("") == []


def test_testo_da_file():
    import docx
    from openpyxl import Workbook
    documento = docx.Document()
    documento.add_paragraph("Clausola di rinnovo tacito annuale.")
    buffer = io.BytesIO()
    documento.save(buffer)
    assert "rinnovo tacito" in conoscenza.testo_da_file("contratto.docx", buffer.getvalue())
    libro = Workbook()
    libro.active.append(["Cliente", "Importo"])
    libro.active.append(["Bianchi", 1200])
    buffer = io.BytesIO()
    libro.save(buffer)
    assert "Bianchi | 1200" in conoscenza.testo_da_file("listino.xlsx", buffer.getvalue())
    assert conoscenza.testo_da_file("nota.txt", "Città: Cagliari".encode()) == "Città: Cagliari"
    assert conoscenza.testo_da_file("rotto.pdf", b"non sono un pdf") == ""
    assert conoscenza.testo_da_file("foto.png", b"\x89PNG") == ""


def test_indice_incrementale(scena, con_servizio):
    primo = conoscenza.indicizza()
    assert primo["nuovi"] >= 4 and primo["vettori"] == primo["nuovi"]               # record e note, con vettore
    nota = Frammento.objects.get(fonte="nota", testo__contains="disdetta")
    assert nota.object_id == scena["bianchi"].pk and "Nota di capo" in nota.titolo
    assert Frammento.objects.get(fonte="record", object_id=scena["bianchi"].pk).testo.startswith("azienda: Bianchi")
    assert conoscenza.indicizza() == {"nuovi": 0, "aggiornati": 0, "tolti": 0, "vettori": 0}   # nulla di nuovo
    Azienda.objects.filter(pk=scena["bianchi"].pk).update(citta="Roma")
    assert conoscenza.indicizza()["aggiornati"] == 1
    scena["bianchi"].delete()
    assert conoscenza.indicizza()["tolti"] == 3                                     # il record e le sue due note
    assert not any(Frammento.objects.filter(testo__contains="password").values_list("pk", flat=True))


def test_indice_senza_servizio_e_ricerca_per_parole(scena):
    assert conoscenza.indicizza()["vettori"] == 0
    risultati = strumenti_conoscenza.cerca_conoscenza(_contesto(scena["capo"], scena["azienda"]),
                                                      "contratto manutenzione Bianchi")
    assert risultati["ricerca"] == "per parole chiave"
    assert risultati["risultati"][0]["titolo"].startswith("Nota di") and "disdetta" in risultati["risultati"][0][
        "estratto"]
    assert risultati["risultati"][0]["url"] == f"/core/aziende/{scena['bianchi'].pk}/"


def test_ricerca_semantica_trova_i_sinonimi(scena, con_servizio):
    conoscenza.indicizza()
    risultati = strumenti_conoscenza.cerca_conoscenza(_contesto(scena["capo"], scena["azienda"]),
                                                      "quando posso fare recesso dal contratto?")
    assert risultati["ricerca"] == "semantica e per parole chiave"
    assert "disdetta" in risultati["risultati"][0]["estratto"] and risultati["risultati"][0]["n"] == 1
    assert ("query", 1) in con_servizio


def test_solo_quello_che_l_utente_puo_vedere(scena, con_servizio):
    conoscenza.indicizza()
    nessuno = User.objects.create_user("nessuno")
    risultati = strumenti_conoscenza.cerca_conoscenza(_contesto(nessuno, scena["azienda"]), "contratto disdetta")
    assert risultati["risultati"] == [] and "Nessun testo" in risultati["nota"]
    nessuno.user_permissions.add(Permission.objects.get(codename="view_azienda"))
    nessuno = User.objects.get(pk=nessuno.pk)
    assert strumenti_conoscenza.cerca_conoscenza(_contesto(nessuno, scena["azienda"]), "contratto disdetta")[
        "risultati"]
    with pytest.raises(strumenti.ErroreStrumento):
        strumenti_conoscenza.cerca_conoscenza(_contesto(nessuno, scena["azienda"]), "  ")


def test_dal_motore_con_citazioni_e_comando(scena, con_servizio, capsys):
    call_command("ai_indicizza")
    assert "Nell'indice:" in capsys.readouterr().out
    provider = Provider.objects.create(nome="Finto", tipo="finto", locale=True)
    modello = Modello.objects.create(provider=provider, id_api="finto-1", etichetta="Finto")
    AnelloCatena.objects.create(profilo=Profilo.objects.create(nome="Chat", compito="chat"), modello=modello)
    conversazione = Conversazione.objects.create(utente=scena["capo"])
    with finto.con_copione([[("strumento", "cerca_conoscenza", {"domanda": "rinnovo contratto Bianchi"})],
                            ["Si rinnova ogni anno a gennaio [1](/core/aziende/2/)."]]) as chiamate:
        motore.avvia(conversazione, "Cosa dice il contratto di Bianchi sul rinnovo?", scena["capo"], scena["azienda"])
    assert "cerca_conoscenza" in chiamate[0]["strumenti"] and "cerca_conoscenza" in chiamate[0]["sistema"]
    risultato = Evento.objects.get(tipo="strumento").dati["risultati"][0]
    assert not risultato["errore"] and '"n": 1' in risultato["contenuto"] and "cita ogni informazione" in risultato[
        "contenuto"]


@con_documenti
def test_documenti_con_le_cartelle_e_lettura(scena, con_servizio, settings, tmp_path):
    from openpyxl import Workbook

    from documenti import servizi
    from documenti.models import Cartella
    settings.PRIVATE_MEDIA_ROOT = tmp_path
    utenti = Group.objects.create(name="Utenti documenti")
    utenti.permissions.set(Permission.objects.filter(content_type__app_label="documenti").exclude(
        codename="vedi_tutte_cartelle"))
    mario = User.objects.create_user("mario", password="x")
    mario.groups.add(utenti)
    pubblica = Cartella.objects.create(azienda=scena["azienda"], nome="Contratti")
    riservata = Cartella.objects.create(azienda=scena["azienda"], nome="Riservata")
    riservata.gruppi_lettura.add(Group.objects.create(name="Capi"))
    testo = ("Contratto Rossi. " * 80 + "Articolo 9: il contratto si proroga di dodici mesi senza comunicazione. ")
    contratto = servizi.carica(pubblica, ContentFile(testo.encode(), name="contratto-rossi.txt"))
    servizi.carica(riservata, ContentFile(b"Stipendi riservati: proroga dei premi.", name="stipendi.txt"))
    libro = Workbook()
    for riga in (["Cliente", "Città", "Fatturato"], ["Rossi", "Cagliari", 1200], ["Bianchi", "Milano", 900]):
        libro.active.append(riga)
    buffer = io.BytesIO()
    libro.save(buffer)
    foglio = servizi.carica(pubblica, ContentFile(buffer.getvalue(), name="clienti.xlsx"))
    conoscenza.indicizza()
    assert Frammento.objects.filter(fonte="documento", object_id=contratto.pk).count() >= 2
    mario = User.objects.get(pk=mario.pk)
    c = _contesto(mario, scena["azienda"])
    trovati = strumenti_conoscenza.cerca_conoscenza(c, "rinnovo del contratto", fonti=["documento"])["risultati"]
    assert trovati and trovati[0]["titolo"].startswith("Documento «contratto-rossi.txt»")
    assert "Articolo 9" in trovati[0]["estratto"] and not any("stipendi" in r["titolo"] for r in trovati)
    pagina = strumenti_conoscenza.leggi_documento(c, contratto.pk)
    assert pagina["pagine"] == 1 and "Articolo 9" in pagina["testo"]
    righe = strumenti_conoscenza.leggi_foglio(c, foglio.pk, testo="cagliari")
    assert righe["colonne"] == ["Cliente", "Città", "Fatturato"] and righe["righe"] == [["Rossi", "Cagliari", 1200]]
    with pytest.raises(strumenti.ErroreStrumento, match="Excel"):
        strumenti_conoscenza.leggi_foglio(c, contratto.pk)


def test_viste_tecniche_fuori_dall_indice(db):
    modelli = {v.opts.label_lower for v, _ in conoscenza._viste_indicizzabili()}
    assert "core.aziende" not in modelli and "core.azienda" in modelli
    assert not modelli & conoscenza.MODELLI_ESCLUSI and not any(m.startswith("ai.") for m in modelli)
