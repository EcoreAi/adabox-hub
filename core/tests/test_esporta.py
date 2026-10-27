import csv
import datetime
import io

import pytest
from django.contrib.auth.models import User
from django.core import mail
from django.utils import timezone
from openpyxl import load_workbook

from core.models import Allegato, Azienda, Lavoro

pytestmark = [pytest.mark.urls("core.tests.urls_viste"), pytest.mark.django_db]


@pytest.fixture(autouse=True)
def privati(settings, tmp_path):
    settings.PRIVATE_MEDIA_ROOT = tmp_path


@pytest.fixture
def lavori(db):
    quando = timezone.make_aware(datetime.datetime(2026, 10, 5, 9, 30))
    return [Lavoro.objects.create(funzione=f, stato=s, tentativi=t, esegui_dopo=quando)
            for f, s, t in [("a.uno", "fatto", 2), ("a.due", "in_coda", 0), ("b.tre", "fallito", 3)]]


def _csv(r):
    testo = r.content.decode("utf-8")
    assert testo.startswith("﻿")
    return list(csv.reader(io.StringIO(testo[1:]), delimiter=";"))


def test_csv_filtrato(client, superuser, lavori):
    r = client.get("/prova/lavori/esporta/", {"formato": "csv", "q": "a."})
    assert r["Content-Type"].startswith("text/csv") and ".csv" in r["Content-Disposition"]
    righe = _csv(r)
    assert righe[0] == ["ID", "Funzione", "Stato", "Tentativi"]
    assert sorted(r_[1:] for r_ in righe[1:]) == [["a.due", "In coda", "0"], ["a.uno", "Fatto", "2"]]


def test_xlsx_con_date(client, superuser, lavori, monkeypatch):
    from core.tests.urls_viste import VISTA
    monkeypatch.setattr(VISTA, "export_fields", ["funzione", "esegui_dopo", "tentativi"])
    r = client.get("/prova/lavori/esporta/", {"formato": "xlsx", "f_stato": "fatto"})
    foglio = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in foglio[1]] == ["Funzione", "Esegui dopo", "Tentativi"]
    assert [c.value for c in foglio[2]] == ["a.uno", datetime.datetime(2026, 10, 5, 9, 30), 2]
    assert foglio.max_row == 2


def test_selezione(client, superuser, lavori):
    r = client.post("/prova/lavori/esporta/", {"formato": "csv", "pk": [lavori[2].pk]})
    assert [r_[1] for r_ in _csv(r)[1:]] == ["b.tre"]


def test_permessi(client, lavori):
    client.force_login(User.objects.create_user("normale", "n@example.com", "x"))
    assert client.get("/prova/lavori/esporta/").status_code == 403


def test_grande_nel_worker(client, superuser, lavori, settings, monkeypatch):
    from core.tests.urls_viste import VISTA
    from core.viste import registry
    monkeypatch.setitem(registry._viste, "prova", [VISTA])      # il lavoro ritrova la vista dal registro
    settings.EXPORT_SINCRONO_MAX = 2
    superuser.email = "root@example.com"
    superuser.save()
    r = client.get("/prova/lavori/esporta/", {"formato": "csv"})
    assert r.status_code == 302
    allegato = Allegato.objects.get()
    assert allegato.record == superuser and allegato.nome.endswith(".csv")
    (email,) = mail.outbox
    assert f"/allegati/{allegato.pk}/" in email.body
    funzioni = {r_[1] for r_ in _csv_da_file(client.get(f"/allegati/{allegato.pk}/"))[1:]}
    assert {"a.uno", "a.due", "b.tre"} <= funzioni        # c'è anche il lavoro di esportazione stesso
    client.force_login(User.objects.create_superuser("altro", "a@example.com", "x"))
    assert client.get(f"/allegati/{allegato.pk}/").status_code == 404


def _csv_da_file(r):
    testo = b"".join(r.streaming_content).decode("utf-8")[1:]
    return list(csv.reader(io.StringIO(testo), delimiter=";"))


def test_bottone_esporta_nella_lista(client, superuser, lavori):
    html = client.get("/prova/lavori/").content.decode()
    assert "/prova/lavori/esporta/?formato=xlsx" in html


def test_csv_decimali_e_date():
    from core import esporta
    dati = esporta.csv_bytes(["a", "b", "c"], [[__import__("decimal").Decimal("1.50"), datetime.date(2026, 1, 2), True]])
    assert dati.decode("utf-8").splitlines()[1] == "1,50;02/01/2026;True"
