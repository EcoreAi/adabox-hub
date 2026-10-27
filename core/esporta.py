"""Export CSV e XLSX delle righe di una ModelView (stessi filtri e permessi della lista)."""
import csv
import datetime
import io
from decimal import Decimal
from types import SimpleNamespace

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.core.mail import send_mail
from django.db import models
from django.http import QueryDict
from django.urls import reverse
from django.utils import timezone

from core.models import Azienda

FORMATI = {"csv": ("text/csv; charset=utf-8", "csv"),
           "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx")}


def campi(vista):
    if vista.export_fields:
        return list(vista.export_fields)
    nomi = {f.name for f in vista.opts.get_fields() if getattr(f, "concrete", False) or f.many_to_many}
    return ["id"] + [c for c in vista.list_display if c in nomi and c != "id"]


def valore(obj, campo):
    """Valore tipizzato per il foglio: numeri e date restano tali, il resto è testo leggibile."""
    f = obj._meta.get_field(campo)
    if f.many_to_many:
        return ", ".join(str(v) for v in getattr(obj, campo).all())
    if f.choices:
        return getattr(obj, f"get_{campo}_display")()
    v = getattr(obj, campo)
    if v is None:
        return ""
    if f.is_relation:
        return str(v)
    if isinstance(v, bool):
        return "Sì" if v else "No"
    if isinstance(v, datetime.datetime):
        return timezone.localtime(v).replace(tzinfo=None) if timezone.is_aware(v) else v
    if isinstance(f, models.FileField):
        return v.name
    return v


def tabella(vista, qs):
    nomi = campi(vista)
    return [vista.etichetta(c) for c in nomi], ([valore(obj, c) for c in nomi] for obj in qs)


def _csv_testo(v):
    if isinstance(v, datetime.datetime):
        return v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, datetime.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, (Decimal, float)):
        return str(v).replace(".", ",")
    return v


def csv_bytes(intestazioni, righe):
    """CSV come lo apre Excel in italiano: separatore ';', virgola decimale, UTF-8 con BOM."""
    testo = io.StringIO()
    scrittore = csv.writer(testo, delimiter=";")
    scrittore.writerow(intestazioni)
    for riga in righe:
        scrittore.writerow([_csv_testo(v) for v in riga])
    return ("﻿" + testo.getvalue()).encode("utf-8")


def xlsx_bytes(intestazioni, righe, titolo="Dati"):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    libro = Workbook(write_only=False)
    foglio = libro.active
    foglio.title = titolo[:31]
    foglio.append(intestazioni)
    for cella in foglio[1]:
        cella.font = Font(bold=True)
    larghezze = [len(str(i)) for i in intestazioni]
    for riga in righe:
        foglio.append(riga)
        for i, v in enumerate(riga):
            larghezze[i] = min(60, max(larghezze[i], len(str(v)) if not isinstance(v, datetime.date) else 12))
        for cella in foglio[foglio.max_row]:
            if isinstance(cella.value, datetime.datetime):
                cella.number_format = "DD/MM/YYYY HH:MM"
            elif isinstance(cella.value, datetime.date):
                cella.number_format = "DD/MM/YYYY"
    for i, w in enumerate(larghezze):
        foglio.column_dimensions[foglio.cell(row=1, column=i + 1).column_letter].width = w + 2
    foglio.freeze_panes = "A2"
    uscita = io.BytesIO()
    libro.save(uscita)
    return uscita.getvalue()


def contenuto(vista, qs, formato):
    intestazioni, righe = tabella(vista, qs)
    if formato == "xlsx":
        return xlsx_bytes(intestazioni, righe, str(vista.verbose_name_plural))
    return csv_bytes(intestazioni, righe)


def nome_file(vista, formato):
    return f"{vista.slug}-{timezone.localtime():%Y%m%d-%H%M}.{FORMATI[formato][1]}"


def queryset(vista, request, pks=None):
    from core.viste import filtri
    qs = filtri.applica(vista, request, vista.get_queryset(request))
    return qs.filter(pk__in=pks) if pks else qs


def in_background(modulo, slug, utente_id, azienda_id, parametri, formato, pks):
    """Lavoro: genera l'export, lo salva come allegato dell'utente e gli manda il link."""
    from core import allegati
    from core.viste import registry
    vista = next(v for v in registry.viste_di(modulo) if v.slug == slug)
    utente = get_user_model().objects.get(pk=utente_id)
    richiesta = SimpleNamespace(user=utente, azienda=Azienda.objects.filter(pk=azienda_id).first(),
                                GET=QueryDict(parametri))
    dati = contenuto(vista, queryset(vista, richiesta, pks), formato)
    allegato = allegati.salva(utente, ContentFile(dati, name=nome_file(vista, formato)), utente)
    link = settings.SITO_URL + reverse("core:allegato_scarica", args=[allegato.pk])
    if utente.email:
        send_mail(f"[{settings.THEME.get('brand', '')}] Esportazione pronta: {vista.verbose_name_plural}",
                  f"L'esportazione che hai chiesto è pronta:\n{link}\n", None, [utente.email])
    return allegato.pk
